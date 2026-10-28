"""Strict parser for the BigSeller "single SKU / all warehouses" inventory export.

Repair v2 (R3) contract §2:
  * headers are matched by **exact normalised alias only** -- no substring fallback.
    A missing required column, a duplicate / ambiguous header blocks the whole file
    with ``HEADER_INVALID``;
  * ``现有库存`` etc. are exact aliases, so an amount column such as
    ``现有库存总成本`` can never be mistaken for a quantity;
  * ``SKU类型/商品类型`` is inspected: combo/combo-set rows become
    ``parseStatus='combo_sku'`` and the file validation becomes ``blocked`` (the whole
    file is not published, even though the row can be attributed to a warehouse);
  * rows are attributed to warehouses using the stable external ids from
    ``warehouse_discovery`` (never fuzzy names); unknown/duplicated names make the
    row ``missing_warehouse`` and block the whole file;
  * the returned payload contains ``fileValidation`` (status/missing/ambiguous/
    duplicate columns/quantityHeader/comboRowCount/unattributableRowCount/reason);
  * amount columns are isolated into the private ``totalCostText`` only;
  * ``rowHash`` keeps the registered field set (including ``warehouseExternalId``),
    because recovery must reproduce exactly the hash stored at registration.

No database or browser code here; unit-testable with a workbook path.
"""
from __future__ import annotations

import hashlib
import json
import re
import warnings
from pathlib import Path
from typing import Any, Iterable

PARSER_VERSION = 'inventory-parser-v2'

warnings.filterwarnings('ignore', message=r'.*no default style.*', category=UserWarning)
warnings.filterwarnings('ignore', message=r'.*default style.*', category=UserWarning,
                        module=r'openpyxl.*')

REQUIRED_FIELDS = ('sku', 'warehouse', 'on_hand', 'order_locked',
                   'warehouse_available', 'in_transit')
REQUIRED_LABELS = {
    'sku': 'SKU编号/SKU信息',
    'warehouse': '仓库',
    'on_hand': '现有库存',
    'order_locked': '订单已锁',
    'warehouse_available': '整仓可用',
    'in_transit': '在途中',
}
QUANTITY_FIELDS = ('on_hand', 'order_locked', 'warehouse_available',
                   'activity_reserved', 'in_transit')
QUANTITY_LABELS = {
    'on_hand': '现有库存',
    'order_locked': '订单已锁',
    'warehouse_available': '整仓可用',
    'activity_reserved': '活动预留',
    'in_transit': '在途中',
}
TEXT_FIELD_KEYS = {
    'on_hand': 'onHandText',
    'order_locked': 'orderLockedText',
    'warehouse_available': 'warehouseAvailableText',
    'activity_reserved': 'activityReservedText',
    'in_transit': 'inTransitText',
}

# Exact aliases after ``normalize_header_name`` (whitespace/case removed).
EXACT_ALIASES: dict[str, tuple[str, ...]] = {
    'sku': ('sku编号', 'sku信息', 'sku编码', 'sku id', 'sku'),
    'warehouse': ('仓库', '仓库名称'),
    'on_hand': ('现有库存', '当前库存', '在手库存', '现有数量', '库存数量'),
    'order_locked': ('订单已锁', '订单锁定', '锁定库存', '锁定数量'),
    'warehouse_available': ('整仓可用', '可用库存', '可用数量'),
    'in_transit': ('在途中', '在途库存', '在途数量'),
    'activity_reserved': ('活动预留', '活动锁定', '预留库存'),
    'total_cost': ('总成本价', '成本价', '总成本', '成本'),
    'weighted_cost': ('加权成本价', '加权成本'),
    'in_transit_cost': ('在途总成本', '在途成本'),
    'gtin': ('gtin', '条码', '条形码', 'barcode'),
    'product_name': ('名称', '商品名称', '产品名称', '商品名'),
    'sku_type': ('sku类型', '商品类型', '产品类型'),
}
AMOUNT_FIELD_ORDER = ('total_cost', 'weighted_cost', 'in_transit_cost')

MAX_HEADER_SCAN_ROWS = 15

_SEPARATOR_RE = re.compile(r'[,\uFF0C\s\u00A0\u202F\u2009\u2007\u2060\uFEFF\'’]')
_INTEGER_RE = re.compile(r'^[+-]?\d+$')
_DECIMAL_INTEGER_RE = re.compile(r'^([+-]?\d+)\.0+$')
_CODE_TOKEN_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._/#+:-]{2,}$')
_WHITESPACE_RE = re.compile(r'\s+')
_MISSING_MARKERS = {'-', '--', '—', '–', '——', '−'}
_COMBO_MARKERS = ('combo', '组合')


class ParseError(RuntimeError):
    """Failure that blocks the whole file (mirrors ExportError shape)."""

    def __init__(self, code: str, detail: str = '', **fields: Any):
        super().__init__(f'{code}: {detail}' if detail else code)
        self.code = code
        self.detail = detail
        self.fields = fields


def _assert_disjoint_aliases() -> None:
    owner: dict[str, str] = {}
    for field, aliases in EXACT_ALIASES.items():
        for alias in aliases:
            if alias in owner and owner[alias] != field:
                raise RuntimeError(f'ambiguous header alias {alias!r}: {owner[alias]} / {field}')
            owner[alias] = field


_assert_disjoint_aliases()


def empty_file_validation() -> dict:
    return {
        'status': 'ok',
        'blocked': False,
        'missingColumns': [],
        'ambiguousColumns': [],
        'duplicateColumns': [],
        'quantityHeader': {},
        'comboRowCount': 0,
        'unattributableRowCount': 0,
        'reason': '',
    }


# ---------------------------------------------------------------------------
# normalisation / cell helpers
# ---------------------------------------------------------------------------

def normalize_header_name(value: Any) -> str:
    text = '' if value is None else str(value)
    text = text.replace('\u3000', ' ').replace('\ufeff', '').strip()
    return _WHITESPACE_RE.sub('', text).lower()


def normalize_warehouse_name(name: str) -> str:
    """Trim/collapse whitespace only -- never fuzzy-match a renamed warehouse."""
    return _WHITESPACE_RE.sub(' ', (name or '').replace('\u3000', ' ').strip()).lower()


def _display_name(value: Any) -> str:
    return '' if value is None else str(value).strip()


def cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    text = str(value).strip()
    return text or None


def parse_quantity(value: Any) -> tuple[int | None, str | None]:
    """Parse a quantity cell. Returns ``(value, invalid_raw)``.

    ``None``/empty/dash -> ``(None, None)`` (unknown is not zero); ``"3,793"`` ->
    ``(3793, None)``; ``"0"``/``"0.0"`` -> ``(0, None)``; anything else ->
    ``(None, raw_text)`` so the row becomes ``invalid_quantity``.
    """
    if value is None:
        return None, None
    if isinstance(value, bool):
        return None, str(value)
    if isinstance(value, int):
        return int(value), None
    if isinstance(value, float):
        if value.is_integer():
            return int(value), None
        return None, cell_text(value) or repr(value)
    text = str(value).strip()
    if text == '' or text in _MISSING_MARKERS:
        return None, None
    cleaned = _SEPARATOR_RE.sub('', text)
    if cleaned == '':
        return None, None
    if _INTEGER_RE.match(cleaned):
        try:
            return int(cleaned), None
        except ValueError:  # pragma: no cover
            return None, text
    decimal_match = _DECIMAL_INTEGER_RE.match(cleaned)
    if decimal_match:
        return int(decimal_match.group(1)), None
    return None, text


def extract_external_sku(raw: Any, *, explicit_code_column: bool = False) -> str | None:
    """Extract the source SKU.

    An explicit ``SKU编号`` column is a real identifier: keep the whole cell string, only trimming
    the ends (internal spaces, punctuation and leading zeros stay). The legacy ``SKU信息`` mix
    column keeps the old token extraction because it can contain label/value text.
    """
    if raw is None:
        return None
    if explicit_code_column:
        text = cell_text(raw)
        if text is None:
            return None
        text = text.replace('\u3000', ' ').strip()
        return text or None
    if not isinstance(raw, str):
        return cell_text(raw)
    text = raw.replace('\u3000', ' ').strip()
    if not text:
        return None
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) > 1:
        tokens = [tok for tok in re.split(r'[\s,，;；]+', lines[0]) if tok]
        for token in tokens:
            if _CODE_TOKEN_RE.match(token):
                return token
        return None
    tokens = [tok for tok in re.split(r'[\s,，;；]+', text) if tok]
    for token in tokens:
        if _CODE_TOKEN_RE.match(token):
            return token
    if len(tokens) == 1:
        return tokens[0]
    return None


def is_combo_type(value: Any) -> bool:
    text = (cell_text(value) or '').lower()
    return bool(text) and any(marker in text for marker in _COMBO_MARKERS)


def _get(mapping: Any, *names: str, default: Any = None) -> Any:
    if isinstance(mapping, dict):
        for name in names:
            if mapping.get(name) is not None:
                return mapping[name]
        return default
    for name in names:
        value = getattr(mapping, name, None)
        if value is not None:
            return value
    return default


def warehouse_discovery_lookup(discovery: Iterable[Any] | None) -> tuple[dict[str, dict], dict[str, list[dict]]]:
    """Return (externalId -> entry, normalised name -> [entries])."""
    by_id: dict[str, dict] = {}
    by_name: dict[str, list[dict]] = {}
    for item in discovery or []:
        external_id = _get(item, 'externalId', 'external_id', 'id')
        name = _get(item, 'name', 'external_name')
        if external_id is None or name is None:
            continue
        entry = {'externalId': str(external_id), 'name': str(name)}
        by_id[entry['externalId']] = entry
        by_name.setdefault(normalize_warehouse_name(entry['name']), []).append(entry)
    return by_id, by_name


# ---------------------------------------------------------------------------
# header detection (exact aliases only)
# ---------------------------------------------------------------------------

def match_headers(header_names: list[Any]) -> dict:
    """Match a candidate header row by exact alias equality (no substring fallback)."""
    normalized = [normalize_header_name(name) for name in header_names]
    display = [_display_name(name) for name in header_names]
    counts: dict[str, int] = {}
    for name in normalized:
        if name:
            counts[name] = counts.get(name, 0) + 1
    duplicate_columns = sorted({display[index] for index, name in enumerate(normalized)
                                if name and counts[name] > 1})
    alias_owner = {alias: field for field, aliases in EXACT_ALIASES.items() for alias in aliases}
    blocking_duplicates = [name for name in duplicate_columns
                           if normalize_header_name(name) in alias_owner]

    mapping: dict[str, int] = {}
    ambiguous: list[str] = []
    missing: list[str] = []
    for field, aliases in EXACT_ALIASES.items():
        hits: list[int] = []
        for index, name in enumerate(normalized):
            if name and name in aliases:
                hits.append(index)
        unique_hits = sorted(set(hits))
        if not unique_hits:
            if field in REQUIRED_FIELDS:
                missing.append(REQUIRED_LABELS[field])
            continue
        if len(unique_hits) > 1:
            ambiguous.append('/'.join(display[index] for index in unique_hits))
            continue
        mapping[field] = unique_hits[0]

    if ambiguous:
        # never keep a partially matched field
        for name in ambiguous:
            for field in list(mapping):
                display_name = display[mapping[field]]
                if display_name in name.split('/'):
                    mapping.pop(field, None)

    amounts: dict[str, int] = {}
    for field in AMOUNT_FIELD_ORDER:
        if field in mapping:
            amounts[field] = mapping[field]
    quantity_header = {field: display[mapping[field]] for field in REQUIRED_FIELDS
                       if field in mapping and field in ('on_hand', 'order_locked',
                                                         'warehouse_available', 'in_transit')}
    return {
        'mapping': mapping,
        'amounts': amounts,
        'missingColumns': missing,
        'ambiguousColumns': ambiguous,
        'duplicateColumns': duplicate_columns,
        'blockingDuplicates': blocking_duplicates,
        'quantityHeader': quantity_header,
    }


def _detect_header(sheet) -> tuple[int, dict]:
    best: tuple[tuple[int, int, int], int, dict] | None = None
    scanned = 0
    for row_index, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN_ROWS, values_only=True),
            start=1):
        scanned = row_index
        names = list(values)
        match = match_headers(names)
        usable = (not match['missingColumns'] and not match['ambiguousColumns']
                  and not match['blockingDuplicates'])
        rank = (1 if usable else 0, len(match['mapping']), -row_index)
        if best is None or rank > best[0]:
            best = (rank, row_index, {**match, 'names': names})
        if usable:
            best = (rank, row_index, {**match, 'names': names})
            break
    if best is None:
        raise ParseError('HEADER_INVALID', '未找到表头行')
    _, header_row, match = best
    missing = match['missingColumns']
    ambiguous = match['ambiguousColumns']
    duplicates = match['blockingDuplicates']
    if missing or ambiguous or duplicates:
        details = []
        if missing:
            details.append('缺少关键列：' + '、'.join(missing))
        if ambiguous:
            details.append('表头歧义（同义列重复）：' + '、'.join(ambiguous))
        if duplicates:
            details.append('表头重名：' + '、'.join(duplicates))
        preview = ' | '.join(_display_name(name) for name in match['names'][:20])
        raise ParseError(
            'HEADER_INVALID',
            '；'.join(details) + f'（扫描到 {scanned} 行；表头候选：{preview or "空"}）',
            missingColumns=missing,
            ambiguousColumns=ambiguous,
            duplicateColumns=duplicates,
        )
    return header_row, match


# ---------------------------------------------------------------------------
# row extraction
# ---------------------------------------------------------------------------

def compute_row_hash(row: dict) -> str:
    """Stable content hash for the DB registered-row completeness check.

    The field set deliberately includes ``warehouseExternalId`` (repair contract §2:
    recovery re-parses with the discovery recorded by the DB, so the same hash must
    be reproducible; discovering is never downgraded).
    """
    canonical = {
        'excelRow': row.get('excelRow'),
        'warehouseExternalId': row.get('warehouseExternalId'),
        'warehouseName': row.get('warehouseName'),
        'externalSku': row.get('externalSku'),
        'productName': row.get('productName'),
        'gtin': row.get('gtin'),
        'onHandText': row.get('onHandText'),
        'orderLockedText': row.get('orderLockedText'),
        'warehouseAvailableText': row.get('warehouseAvailableText'),
        'activityReservedText': row.get('activityReservedText'),
        'inTransitText': row.get('inTransitText'),
        'parseStatus': row.get('parseStatus'),
    }
    serialized = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return hashlib.sha256(serialized.encode('utf-8')).hexdigest()


def _raw_cell(row_values: tuple, index: int | None) -> Any:
    if index is None or index >= len(row_values):
        return None
    return row_values[index]


def _total_cost_text(row_values: tuple, amounts: dict[str, int]) -> str | None:
    for field in AMOUNT_FIELD_ORDER:
        text = cell_text(_raw_cell(row_values, amounts.get(field)))
        if text and text not in _MISSING_MARKERS:
            return text
    return None


def _attribute_warehouse(warehouse_name: str | None,
                         by_name: dict[str, list[dict]]) -> tuple[dict | None, str | None]:
    if not warehouse_name:
        return None, '仓库列为空'
    candidates = by_name.get(normalize_warehouse_name(warehouse_name)) or []
    if len(candidates) == 1:
        return candidates[0], None
    if len(candidates) > 1:
        return None, f'仓库名称“{warehouse_name}”对应多个外部仓库 id'
    return None, f'仓库名称“{warehouse_name}”不在本次导出发现中'


def _row_from_values(row_values: tuple, excel_row: int, match: dict,
                     by_name: dict[str, list[dict]]) -> tuple[dict, str | None]:
    mapping = match['mapping']
    amounts = match['amounts']
    sku_index = mapping.get('sku')
    sku_header = normalize_header_name(match['names'][sku_index]) if sku_index is not None and sku_index < len(match.get('names') or []) else ''
    # Only the explicit ``SKU编号`` header is a literal identifier; legacy mixed columns keep parsing.
    sku = extract_external_sku(_raw_cell(row_values, sku_index), explicit_code_column=(sku_header == 'sku编号'))
    warehouse_name = cell_text(_raw_cell(row_values, mapping.get('warehouse')))
    discovery, problem = _attribute_warehouse(warehouse_name, by_name)
    external_id = discovery['externalId'] if discovery else None

    row: dict[str, Any] = {
        'excelRow': excel_row,
        'warehouseExternalId': external_id,
        'warehouseName': warehouse_name,
        'externalSku': sku,
        'productName': cell_text(_raw_cell(row_values, mapping.get('product_name'))),
        'gtin': cell_text(_raw_cell(row_values, mapping.get('gtin'))),
        'onHandText': None,
        'orderLockedText': None,
        'warehouseAvailableText': None,
        'activityReservedText': None,
        'inTransitText': None,
        'totalCostText': _total_cost_text(row_values, amounts),
        'rowHash': '',
        'parseStatus': 'ok',
        'reason': None,
    }

    invalid_parts: list[str] = []
    for field in QUANTITY_FIELDS:
        raw_value = _raw_cell(row_values, mapping.get(field))
        row[TEXT_FIELD_KEYS[field]] = cell_text(raw_value)
        _parsed, invalid_raw = parse_quantity(raw_value)
        if invalid_raw is not None:
            invalid_parts.append(f'{QUANTITY_LABELS[field]}="{invalid_raw}"')

    combo = is_combo_type(_raw_cell(row_values, mapping.get('sku_type')))
    if problem:
        row['parseStatus'] = 'missing_warehouse'
        row['reason'] = f'{problem}，无法归属仓库'
    elif combo:
        row['parseStatus'] = 'combo_sku'
        row['reason'] = '来源 SKU 类型为组合/Combo，整份文件不采用'
    elif not sku:
        row['parseStatus'] = 'missing_sku'
        row['reason'] = 'SKU 列为空，无法对应商品'
    elif invalid_parts:
        row['parseStatus'] = 'invalid_quantity'
        row['reason'] = '数量无法解析：' + '，'.join(invalid_parts)
    row['rowHash'] = compute_row_hash(row)
    return row, (problem if row['parseStatus'] == 'missing_warehouse' else None)


def _mark_duplicates(rows: list[dict]) -> None:
    seen: dict[tuple[str, str], list[dict]] = {}
    for row in rows:
        if row.get('warehouseName') and row.get('externalSku'):
            key = (normalize_warehouse_name(row['warehouseName']), row['externalSku'])
            seen.setdefault(key, []).append(row)
    for group in seen.values():
        if len(group) < 2:
            continue
        excel_rows = [int(row['excelRow']) for row in group]
        reason = (f'同仓同SKU重复 {len(group)} 次（Excel行 '
                  f'{",".join(str(number) for number in excel_rows)}），未自动采用任一行')
        for row in group:
            if row.get('parseStatus') == 'ok':
                row['parseStatus'] = 'duplicate_sku'
                row['reason'] = reason


# ---------------------------------------------------------------------------
# public entry points
# ---------------------------------------------------------------------------

def parse_export_file(
    path: str | Path,
    *,
    warehouse_discovery: Iterable[Any] | None = None,
    expected_sha256: str | None = None,
    storage_relpath: str | None = None,
    export_mark: str | None = None,
    export_started_at: str | None = None,
    export_finished_at: str | None = None,
) -> dict:
    """Parse one export workbook into register-payload pieces (contract §2)."""
    file_path = Path(path)
    if not file_path.is_file():
        raise ParseError('FILE_MISSING', f'文件不存在: {file_path.name}')
    if file_path.stat().st_size <= 0:
        raise ParseError('FILE_EMPTY', f'文件为空: {file_path.name}')

    actual_sha = _sha256_file(file_path)
    if expected_sha256 and actual_sha != expected_sha256:
        raise ParseError('FILE_CORRUPT', '文件 sha256 与导出登记不一致')

    discovery_entries = list(warehouse_discovery or [])
    if not discovery_entries:
        raise ParseError('INVENTORY_WAREHOUSE_DISCOVERY_MISSING',
                         '解析需要本次导出的仓库发现结果（稳定 externalId + name）')

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError('PARSER_UNAVAILABLE', 'openpyxl 未安装') from exc

    # BigSeller writes a wrong worksheet <dimension> (e.g. A1 while the sheet really
    # has 28 columns). read_only mode trusts it and would read only the first column,
    # so normal mode MUST be used (files are small).
    try:
        workbook = load_workbook(filename=str(file_path), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError('FILE_CORRUPT', f'无法打开工作簿（{type(exc).__name__}）') from None

    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_row, match = _detect_header(sheet)
        header_values = match['names']
        header = [{'index': index, 'name': _display_name(name)}
                  for index, name in enumerate(header_values)]
        while header and not header[-1]['name']:
            header.pop()
        _by_id, by_name = warehouse_discovery_lookup(discovery_entries)

        rows: list[dict] = []
        blank_rows = 0
        unattributable = 0
        for offset, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1):
            if _is_blank_row(values):
                blank_rows += 1
                continue
            row, problem = _row_from_values(values, offset, match, by_name)
            if row['parseStatus'] == 'missing_warehouse':
                unattributable += 1
            rows.append(row)
        if not rows:
            raise ParseError('FILE_EMPTY', '表头存在但没有任何数据行')

        _mark_duplicates(rows)
        for row in rows:
            row['rowHash'] = compute_row_hash(row)

        combo_rows = sum(1 for row in rows if row['parseStatus'] == 'combo_sku')
        validation = empty_file_validation()
        validation.update({
            'missingColumns': match['missingColumns'],
            'ambiguousColumns': match['ambiguousColumns'],
            'duplicateColumns': match['duplicateColumns'],
            'quantityHeader': match['quantityHeader'],
            'comboRowCount': combo_rows,
            'unattributableRowCount': unattributable,
        })
        reasons = []
        if combo_rows:
            reasons.append(f'发现 {combo_rows} 行组合/Combo SKU，整份文件不发布')
        if unattributable:
            reasons.append(f'发现 {unattributable} 行无法归属仓库，整份文件不发布')
        if reasons:
            validation['status'] = 'blocked'
            validation['blocked'] = True
            validation['reason'] = '；'.join(reasons)

        file_info = {
            'storageRelpath': storage_relpath or str(file_path),
            'sha256': actual_sha,
            'sizeBytes': file_path.stat().st_size,
            'sheetName': sheet.title,
            'header': header,
            'rowCount': len(rows),
            'blankRowCount': blank_rows,
        }
        return {
            'exportMark': export_mark,
            'exportStartedAt': export_started_at,
            'exportFinishedAt': export_finished_at,
            'file': file_info,
            'rows': rows,
            'blankRowCount': blank_rows,
            'fileValidation': validation,
            'headerRow': header_row,
            'headerMapping': {
                field: _display_name(header_values[index]) if index < len(header_values) else ''
                for field, index in match['mapping'].items()
            },
            'parserVersion': PARSER_VERSION,
        }
    finally:
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            pass


parse_export_workbook = parse_export_file


def parse_adapter_rows(rows: Iterable[Any], *, blank_row_count: int = 0,
                       warehouse_discovery: Iterable[Any] | None = None) -> tuple[list[dict], dict]:
    """Fallback for an adapter that already returns rows (the real adapter returns a file)."""
    _by_id, by_name = warehouse_discovery_lookup(warehouse_discovery)
    normalized: list[dict] = []
    for raw in rows or []:
        if not isinstance(raw, dict):
            raise ParseError('INVALID_EXPORT_PAYLOAD', '导出行必须是对象')
        warehouse_name = raw.get('warehouseName')
        external_id = raw.get('warehouseExternalId')
        if external_id is None and warehouse_name:
            discovery, _problem = _attribute_warehouse(str(warehouse_name), by_name)
            external_id = discovery['externalId'] if discovery else None
        row = {
            'excelRow': raw.get('excelRow'),
            'warehouseExternalId': external_id,
            'warehouseName': warehouse_name,
            'externalSku': raw.get('externalSku'),
            'productName': raw.get('productName'),
            'gtin': raw.get('gtin'),
            'onHandText': raw.get('onHandText'),
            'orderLockedText': raw.get('orderLockedText'),
            'warehouseAvailableText': raw.get('warehouseAvailableText'),
            'activityReservedText': raw.get('activityReservedText'),
            'inTransitText': raw.get('inTransitText'),
            'totalCostText': raw.get('totalCostText'),
            'rowHash': raw.get('rowHash'),
            'parseStatus': raw.get('parseStatus') or 'ok',
            'reason': raw.get('reason'),
        }
        if not isinstance(row['excelRow'], int):
            raise ParseError('INVALID_EXPORT_PAYLOAD', '导出行缺少 excelRow')
        normalized.append(row)
    if not normalized:
        raise ParseError('FILE_EMPTY', f'导出适配器未提供数据行（空行 {blank_row_count}）')
    _mark_duplicates(normalized)
    for row in normalized:
        row['rowHash'] = compute_row_hash(row)
    validation = empty_file_validation()
    combo_rows = sum(1 for row in normalized if row['parseStatus'] == 'combo_sku')
    missing_warehouse = sum(1 for row in normalized if row['parseStatus'] == 'missing_warehouse')
    validation['comboRowCount'] = combo_rows
    validation['unattributableRowCount'] = missing_warehouse
    if combo_rows or missing_warehouse:
        validation['status'] = 'blocked'
        validation['blocked'] = True
        validation['reason'] = '导出行包含组合 SKU 或无法归属仓库的行，整份文件不发布'
    return normalized, validation


def find_export_file(export_result: dict, *, work_dir: str | Path,
                     export_dir: str | Path | None = None) -> Path | None:
    file_info = export_result.get('file') or {}
    candidates_raw = [
        export_result.get('filePath'),
        file_info.get('filePath') if isinstance(file_info, dict) else None,
        file_info.get('storageRelpath') if isinstance(file_info, dict) else None,
    ]
    work_root = Path(work_dir)
    export_root = Path(export_dir) if export_dir else work_root
    for raw in candidates_raw:
        if not raw:
            continue
        candidate = Path(str(raw))
        options = [candidate] if candidate.is_absolute() else [
            work_root / candidate, export_root / candidate, candidate,
        ]
        for option in options:
            try:
                if option.is_file():
                    return option
            except OSError:  # pragma: no cover
                continue
    return None


def storage_relpath_for(path: str | Path, work_dir: str | Path) -> str:
    file_path = Path(path)
    try:
        return file_path.resolve().relative_to(Path(work_dir).resolve()).as_posix()
    except (ValueError, OSError):
        return str(file_path)


def _is_blank_row(values: Iterable[Any]) -> bool:
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return False
    return True


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def summarize_parse(rows: list[dict]) -> dict:
    counts: dict[str, int] = {}
    for row in rows:
        status = row.get('parseStatus') or 'unknown'
        counts[status] = counts.get(status, 0) + 1
    return counts
