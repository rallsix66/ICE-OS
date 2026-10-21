"""Finite final sample decisions; no matching, external writes or permanent mapping."""
from __future__ import annotations

import copy
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from .models import MasterRow
from .review_exports import export_effective_outputs, refresh_confirmation_summary, review_records, write_csv
from .review_result import MANUAL_FIELDS, read_decisions, resolve_reviews, source_identity

BASE_HASHES = {
    'match-bundle.json': '3c2babc9b13ca11fab14cb4bbce5e8617052319d06bc80c922256f6893346159',
    'manual-decisions.effective.csv': '48544cd03154131970fce220fb2ff7c704e18d7860e2ea7c72e6a0ef1ac8d921',
}
AUTHORIZED = (
    ('mc-e1b56fb81b0db46c', 376, 19, 'WM0012', 'confirm', 'version_reaffirmation'),
    ('mc-73ac4812d7d8de64', 380, 19, 'WM0012', 'confirm', 'version_reaffirmation'),
    ('mc-b7bd376e964738cd', 526, 19, 'WM0012', 'confirm', 'version_reaffirmation'),
    ('mc-e78031a3b38c39f1', 743, 13, 'WM0006', 'ignore', 'batch_ignore'),
    ('mc-06e9251af096ae18', 745, 14, 'WM0008', 'ignore', 'batch_ignore'),
    ('mc-75e0ce710d064c4c', 857, 121, 'WM0130', 'confirm', 'new_selection'),
    ('mc-44b8f0236d346bc7', 220, 67, 'WM0073', 'confirm', 'resolve_defer'),
    ('mc-c8e6e17fc25361f1', 321, 67, 'WM0073', 'confirm', 'resolve_defer'),
)
SPECS = (('WM0006', 13, None), ('WM0008', 14, None), ('WM0031', 25, '20ml/支'),
         ('WM0062', 83, '20ml/支'), ('WM0098', 92, '120g/瓶'), ('WM0130', 121, '20ml/支'))


def final_field_changes(before, after, path=''):
    """Reconstructable JSON Pointer paths, including recursive evidence filenames."""
    if before == after:
        return []
    if isinstance(before, dict) and isinstance(after, dict):
        changes = []
        for key in sorted(set(before) | set(after)):
            escaped_key = str(key).replace('~', '~0').replace('/', '~1')
            location = path+'/'+escaped_key
            if key not in before or key not in after:
                changes.append({'path': location, 'before_present': key in before, 'after_present': key in after,
                                'before': before.get(key), 'after': after.get(key)})
            else:
                changes.extend(final_field_changes(before[key], after[key], location))
        return changes
    if isinstance(before, list) and isinstance(after, list) and len(before) == len(after):
        return [change for i, (a, b) in enumerate(zip(before, after))
                for change in final_field_changes(a, b, path+'/'+str(i))]
    return [{'path': path, 'before': before, 'after': after}]


def validate_plan(base, plan):
    rows = [(p.get('candidate_id'), (p.get('source_file') or {}).get('excel_row'), p.get('target_row'),
             p.get('target_ice'), p.get('decision'), p.get('operation')) for p in plan.get('source_decisions', [])]
    if rows != list(AUTHORIZED):
        raise ValueError('Final decisions do not equal the eight authorized source operations')
    if base.get('final_confirmation_events'):
        raise ValueError('Already-applied final output is not the approved input baseline')
    if len({p.get('decision_id') for p in plan['source_decisions']}) != len(rows) or any(
            not p.get('reason') or not p.get('decision_id') for p in plan['source_decisions']):
        raise ValueError('Final decisions need unique operation IDs and explicit reasons')
    versions = plan.get('version_confirmations', [])
    if [(v.get('candidate_id'), v.get('target_row'), v.get('target_ice'), v.get('confirmed_version'))
            for v in versions] != [(p[0], 19, 'WM0012', '升级普通版') for p in AUTHORIZED[:3]]:
        raise ValueError('Only the three authorized upgraded cleanser versions may be confirmed')
    for v, p in zip(versions, plan['source_decisions']):
        if v.get('decision_id') != p['decision_id'] or v.get('reason') != p['reason']:
            raise ValueError('Version confirmation lacks the corresponding explicit source decision')
    resolutions = plan.get('specification_resolutions', [])
    if [(r.get('ice_code'), r.get('master_row'), r.get('selected_spec')) for r in resolutions] != list(SPECS):
        raise ValueError('Specification decisions do not match the six authorized products')
    for r in resolutions:
        expected = 'closed_ignored_for_batch' if r['selected_spec'] is None else 'closed_master_column_c'
        if (r.get('status') != expected or r.get('source_column') != (None if r['selected_spec'] is None else 'C')
                or r.get('quantity_conversion_authorized') is not False or not r.get('reason')
                or (r.get('previous_audit_question') or {}).get('id') != 'audit-spec-'+r['ice_code']):
            raise ValueError('Specification closure must retain evidence without unit conversion')
    brushes = plan.get('brush_preservation', [])
    if [(p.get('ice_code'), p.get('master_row')) for p in brushes] != [('WM0081', 62), ('brush', 107)] or any(
            p.get('status') != 'pending_confirmation' or p.get('do_not_merge') is not True
            or not p.get('future_requirement') for p in brushes):
        raise ValueError('Both separate brush products must remain pending confirmation')


def _source_ids(candidate):
    return {candidate['candidate_id'], *candidate.get('candidate_id_lineage', []),
            *(r['candidate_id'] for r in candidate.get('decision_records', []))}


def apply_final_decisions(base, decisions, plan):
    validate_plan(base, plan)
    bundle, effective = copy.deepcopy(base), copy.deepcopy(decisions)
    by_id = {c['candidate_id']: c for c in bundle['candidates']}
    masters = {m['row_index']: m for m in bundle['master_rows']}
    if len(by_id) != len(bundle['candidates']):
        raise ValueError('Duplicate candidate ID in approved baseline')
    prior_pending = copy.deepcopy(bundle['batch_business_pending'])
    expected_ids = [p[0] for p in AUTHORIZED[:3]]
    if (len(prior_pending) != 1 or prior_pending[0]['number'] != 1
            or prior_pending[0]['candidate_ids'] != expected_ids):
        raise ValueError('The expected cleanser version issue is not the sole current business issue')
    events = []
    for p in plan['source_decisions']:
        cid = p['candidate_id']
        c, target = by_id.get(cid), masters.get(p['target_row'])
        if (not c or tuple(p['source_identity']) != source_identity(c) or c['source'] != p['source']
                or not target or target != p['target_master'] or target['ice_code'] != p['target_ice']):
            raise ValueError('Final source identity or selected master target differs from frozen evidence')
        source, file = c['source'], p['source_file']
        row_key = source['source_row_key'].split(':')
        original_row = int(row_key[1] if source['source_system'] == 'jushuitan' else row_key[2])
        if original_row != file['excel_row'] or source['sheet_name'] != file['sheet_name']:
            raise ValueError('Final source original row or sheet mismatch')
        expected = {'version_reaffirmation': ('manual_selected', 'confirm'), 'batch_ignore': ('candidate', None),
                    'new_selection': ('candidate', None), 'resolve_defer': ('pending', 'defer')}[p['operation']]
        if (c['review_state'], c.get('review_decision')) != expected:
            raise ValueError('This final authorization cannot override an unexpected current decision')
        if p['operation'] == 'version_reaffirmation' and c.get('selected_master') != target:
            raise ValueError('Version reaffirmation may not replace the existing product selection')
        ids = _source_ids(c)
        if any(other['candidate_id'] != cid and ids & _source_ids(other) for other in by_id.values()):
            raise ValueError('Ambiguous frozen source ID lineage')
        previous = [copy.deepcopy(effective[i]) for i in sorted(ids) if i in effective]
        signatures = {tuple(r.get(k, '') for k in ('review_decision', 'review_notes', *MANUAL_FIELDS)) for r in previous}
        if not previous or len(signatures) != 1:
            raise ValueError('Missing or conflicting inherited CSV decisions for frozen source')
        raw = copy.deepcopy(previous[0])
        if raw.get('review_decision', '') != (expected[1] or '') or raw.get('manual_jst_candidate_id') or raw.get('manual_jst_sku'):
            raise ValueError('Inherited CSV decision or JST reference contradicts this authorization')
        for old_id in ids:
            effective.pop(old_id, None)
        raw.update(candidate_id=cid, review_decision=p['decision'], review_notes=p['reason'],
                   manual_master_row=str(target['row_index']) if p['decision'] == 'confirm' else '',
                   manual_master_ice=p['target_ice'] if p['decision'] == 'confirm' else '')
        effective[cid] = raw
        events.append({'decision_id': p['decision_id'], 'type': 'final_source_decision',
            'explicit_decision': copy.deepcopy(p), 'before_csv_records': previous, 'after_csv': copy.deepcopy(raw),
            'before_selected_master': copy.deepcopy(c.get('selected_master')),
            'before_decision_records': copy.deepcopy(c.get('decision_records', [])),
            'before_batch_business_pending': copy.deepcopy(c.get('batch_business_pending'))})
        if p['operation'] == 'resolve_defer':
            bundle.setdefault('batch_target_authorizations', []).append({
                'candidate_id': cid, 'source_identity': copy.deepcopy(p['source_identity']), 'source': copy.deepcopy(source),
                'decision': 'confirm', 'target_row': p['target_row'], 'target_ice': p['target_ice'],
                'reason': p['reason'], 'authorization': plan['authorization'], 'group': 'final_sample_record_only'})
    bundle['batch_business_pending'] = []
    for cid in expected_ids:
        bundle['candidates'][next(i for i, c in enumerate(bundle['candidates']) if c['candidate_id'] == cid)]['batch_business_pending'] = None
    bundle['resolved_business_pending_history'] = [{'previous_pending': prior_pending[0],
        'resolution_decision_ids': [p['decision_id'] for p in plan['source_decisions'][:3]],
        'reason': '用户明确三条均为升级普通版WM0012；解除当前版本待决，保留此前待决历史。'}]
    outcome = resolve_reviews(bundle, effective, [MasterRow(**m) for m in bundle['master_rows']])
    acted = {p['candidate_id'] for p in plan['source_decisions']}
    for c in outcome['candidates']:
        if c['candidate_id'] not in acted:
            c['previous_review'] = copy.deepcopy(by_id[c['candidate_id']].get('previous_review'))
            if c != by_id[c['candidate_id']]:
                raise ValueError('Final decisions would change an unauthorized source record')
    after = {c['candidate_id']: c for c in outcome['candidates']}
    for p in plan['source_decisions']:
        c = after[p['candidate_id']]
        if (p['decision'] == 'confirm' and (c['review_state'] != 'manual_selected' or c['selected_master'] != p['target_master'])):
            raise ValueError('An authorized final confirmation did not become an effective selection')
        if p['decision'] == 'ignore' and (c['review_state'] != 'ignored' or c['selected_master'] is not None):
            raise ValueError('A final ignore must leave an empty effective target')
    for c in base['candidates']:
        if c.get('selected_master') and after[c['candidate_id']].get('selected_master') != c['selected_master']:
            raise ValueError('An existing effective product target changed')
    outcome['final_confirmation_events'] = events
    outcome['batch_version_confirmations'] = copy.deepcopy(plan['version_confirmations'])
    outcome['batch_specification_resolutions'] = copy.deepcopy(plan['specification_resolutions'])
    outcome['batch_brush_preservation'] = copy.deepcopy(plan['brush_preservation'])
    for r in outcome['batch_specification_resolutions']:
        detail = next(m for m in outcome['master_provenance'] if m['row_index'] == r['master_row'])
        if detail['ice_code'] != r['ice_code'] or (r['selected_spec'] is not None and
                (detail['data']['规格'] != r['selected_spec'] or detail['spec'] != r['selected_spec'])):
            raise ValueError('Approved column C specification differs from preserved master data')
        detail['specification_resolution'] = copy.deepcopy(r)
    for p in outcome['batch_brush_preservation']:
        detail = next(m for m in outcome['master_provenance'] if m['row_index'] == p['master_row'])
        if detail['ice_code'] != p['ice_code'] or detail['name_cn'] != p['name']:
            raise ValueError('Separate brush identity differs from approved master data')
        detail['brush_preservation'] = copy.deepcopy(p)
    refresh_confirmation_summary(outcome)
    return outcome, effective, events


def validate_source_files(base, bundle, plan):
    """Validate exact original input hashes, selected raw rows, and C-column evidence."""
    import openpyxl
    inputs = bundle['input_manifest']['original_input_hashes']
    if plan.get('original_inputs') != inputs:
        raise ValueError('Original input provenance does not match the approved baseline')
    all_inputs = [inputs['master'], inputs['bigseller'], *inputs['jushuitan']['paths']]
    for item in all_inputs:
        if hashlib.sha256((base/item['frozen']).read_bytes()).hexdigest() != item['sha256']:
            raise ValueError('Frozen original input hash mismatch: '+item['frozen'])
    grouped = {}
    for p in plan['source_decisions']:
        file = p['source_file']
        original = next((i for i in all_inputs if i['frozen'] == file['frozen']), None)
        if not original or file['file_name'] != Path(original['path']).name or file['sha256'] != original['sha256']:
            raise ValueError('Final decision original source file or hash mismatch')
        grouped.setdefault((file['frozen'], file['sheet_name']), {})[file['excel_row']] = p
    for (name, sheet_name), wanted in grouped.items():
        book = openpyxl.load_workbook(base/name, read_only=True, data_only=False)
        try:
            sheet = book[sheet_name]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)
            found = set()
            for n, values in enumerate(rows, 2):
                if n in wanted:
                    p = wanted[n]
                    raw = {str(k): (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in zip(headers, values)}
                    if raw != p['source_file']['raw_record']:
                        raise ValueError('Final raw source row differs from frozen Excel')
                    if p['source']['source_system'] == 'jushuitan':
                        if (raw.get('商品编码') != p['source']['external_sku_raw'] or
                                raw.get('商品名称') != p['source']['external_name']):
                            raise ValueError('Final JST source does not agree with its original row')
                    found.add(n)
                if n >= max(wanted):
                    break
            if found != set(wanted):
                raise ValueError('An authorized raw source row does not exist')
        finally:
            book.close()
    old_questions = base/plan['previous_audit_questions']['file']
    if hashlib.sha256(old_questions.read_bytes()).hexdigest() != plan['previous_audit_questions']['sha256']:
        raise ValueError('Previous specification audit evidence hash mismatch')
    questions = json.loads(old_questions.read_text(encoding='utf-8'))
    for r in plan['specification_resolutions']:
        if r['previous_audit_question'] != next(q for q in questions if q['id'] == 'audit-spec-'+r['ice_code']):
            raise ValueError('Previous specification question evidence was changed')
    book = openpyxl.load_workbook(base/inputs['master']['frozen'], read_only=True, data_only=False)
    try:
        sheet = book['商品库']
        for r in plan['specification_resolutions']:
            if r['source_column'] == 'C' and sheet.cell(r['master_row'], 3).value != r['selected_spec']:
                raise ValueError('Explicit selected specification is not the frozen master C cell')
    finally:
        book.close()


def cmd_finalize_sample(args):
    from .cli import _freeze, _json, _new_output, _sha256
    base, plan_path = Path(args.evidence).resolve(), Path(args.decisions).resolve()
    plan = _json(plan_path)
    if plan.get('base_hashes') != BASE_HASHES:
        raise ValueError('Final decision pins do not identify the approved JST baseline')
    for name, digest in BASE_HASHES.items():
        if _sha256(base/name) != digest:
            raise ValueError('Approved final input hash mismatch: '+name)
    old = _json(base/'match-bundle.json')
    validate_source_files(base, old, plan)
    outcome, effective, events = apply_final_decisions(old, read_decisions(base/'manual-decisions.effective.csv'), plan)
    out = _new_output(args.output)
    prior_files = sorted(p for p in base.rglob('*') if p.is_file())
    manifest = {'scope': plan['scope'], 'generated_at': plan['recorded_at'], 'base_evidence': str(base),
        'base_evidence_hashes': {p.relative_to(base).as_posix(): _sha256(p) for p in prior_files},
        'baseline_bundle': _freeze(base/'match-bundle.json', out, 'baseline-match-bundle.frozen.json'),
        'decision_source': _freeze(base/'manual-decisions.effective.csv', out, 'baseline-effective-decisions.frozen.csv'),
        'explicit_decisions': _freeze(plan_path, out, 'final-decisions.frozen.json'),
        'baseline_manifest': _freeze(base/'input-manifest.json', out, 'baseline-input-manifest.frozen.json'),
        'original_input_hashes': copy.deepcopy(old['input_manifest']['original_input_hashes'])}
    inherited = out/'inherited-baseline'
    inherited.mkdir()
    for path in prior_files:
        dest = inherited/path.relative_to(base)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(path.read_bytes())
    root_inherited = ('master-row-lineage.json', 'source-formulas-and-errors.json', 'explicit-correction-audit.json',
        'business-corrections.frozen.json', 'business-decisions.frozen.json', 'jst-decisions.frozen.json',
        'original-22-comparison.json', 'manual-decisions.original.frozen.csv', 'input-0.frozen.xlsx',
        'input-1.frozen.xlsx', 'input-2.frozen.xlsx', 'input-3.frozen.xlsx', 'history-bundle-0.frozen.json', 'history-bundle-1.frozen.json')
    manifest['inherited_files'] = [_freeze(base/name, out, name) for name in root_inherited]
    manifest['complete_inherited_baseline'] = 'inherited-baseline/'
    write_csv(out/'manual-decisions.effective.csv', list(effective.values()), ['candidate_id', 'review_decision', 'review_notes', *MANUAL_FIELDS])
    manifest['effective_decision_source'] = {'frozen': 'manual-decisions.effective.csv',
        'sha256': _sha256(out/'manual-decisions.effective.csv'), 'derived_from': ['decision_source', 'explicit_decisions']}
    outcome['input_manifest'] = manifest
    export_effective_outputs(outcome, out)
    before, after = ({c['candidate_id']: c for c in b['candidates']} for b in (old, outcome))
    comparisons = [{'candidate_id': p['candidate_id'], 'decision_id': p['decision_id'], 'operation': p['operation'],
        'before': before[p['candidate_id']], 'after': after[p['candidate_id']], 'reason': p['reason']} for p in plan['source_decisions']]
    old_rows = {r['candidate_id']: r for r in review_records(old)}
    new_rows = review_records(outcome)
    changes = [{'candidate_id': r['candidate_id'], 'before': old_rows[r['candidate_id']], 'after': r,
        'changed_fields': [k for k in r if r[k] != old_rows[r['candidate_id']].get(k)]}
        for r in new_rows if any(r.get(k) != old_rows[r['candidate_id']].get(k) for k in set(r) | set(old_rows[r['candidate_id']]))]
    old_masters = {m['row_index']: m for m in old['master_provenance']}
    master_changes = [{'row_index': m['row_index'], 'before': old_masters[m['row_index']], 'after': m}
        for m in outcome['master_provenance'] if m != old_masters[m['row_index']]]
    flat = [{'candidate_id': p['candidate_id'], 'operation': p['operation'], 'source_row_key': p['source']['source_row_key'],
        'before_decision': before[p['candidate_id']].get('review_decision') or '',
        'after_decision': after[p['candidate_id']].get('review_decision') or '',
        'before_target': json.dumps(before[p['candidate_id']].get('selected_master'), ensure_ascii=False),
        'after_target': json.dumps(after[p['candidate_id']].get('selected_master'), ensure_ascii=False), 'reason': p['reason']}
        for p in plan['source_decisions']]
    write_csv(out/'final-source-comparison.csv', flat)
    values = [('input-manifest.json', manifest), ('current-final-decision-audit.json', events),
        ('final-source-comparison.json', comparisons), ('all-correspondence-changes.json', changes),
        ('all-master-changes.json', master_changes), ('all-field-changes.json', final_field_changes(old, outcome)),
        ('current-original-22-comparison.json', [{'candidate_id': r['candidate_id'], 'before': before[r['candidate_id']],
            'after': after[r['candidate_id']]} for r in _json(base/'current-original-22-comparison.json')])]
    for name, value in values:
        (out/name).write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps({'review': outcome['review_summary'], 'final_decisions': outcome['final_confirmation_summary']}, ensure_ascii=False))
    return 0
